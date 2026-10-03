"""Pin answer-key handling in the importers (2026-10-03 fix step 5a, audit A2-03/A2-06).

1. A trailing answer key ('빠른 정답 1. ③ 2. ⑤', '정답 및 해설 1번 정답 ②', or a header
   after the exam title such as '2025학년도 … 화학I 정답 및 해설 / [정답표] 1. ③') used to
   become fake duplicate-number questions in TXT/DOCX/HWPX/web/PDF imports. It is
   now cut before question splitting and its answers move to the matching
   question's answer field, with a notice.
2. HWPX answer endnotes that miss the endnote-boundary conditions (1..N, 80%
   answers) used to be silently dropped. They now attach to the question that
   holds the endnote anchor, with a notice.
3. Questions with inline '[정답] ③' lines (PDF page, soft-broken DOCX paragraph) and a stem
   line ending in '정답표' before numeric choices are not an answer section (repair R1).
4. /api/collect always failed with 400 'not enough values to unpack' because the
   collector passed 2-tuples to the 3-tuple chunker.

Mostly synthetic inputs (the official CSAT answer-key check reads data/ originals and
SKIPs only that part when they are absent; a local http.server serves the web page; the SSRF guard
is bypassed for that loopback fetch inside this process only).
Exit codes: 0 = PASS, 1 = FAIL.
"""
from __future__ import annotations

import base64
import functools
import http.server
import io
import os
from pathlib import Path
import sys
import tempfile
import threading
import zipfile

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
_TMP = tempfile.TemporaryDirectory(prefix="hwpmake_answer_key_", ignore_cleanup_errors=True)
os.environ["HWP_MAKE_DATA_DIR"] = _TMP.name

import fitz  # noqa: E402
from docx import Document  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

from app import collector, importers, main, storage  # noqa: E402

failures: list[str] = []
QUESTIONS = "1. 첫째 문항입니다.\n① 가 ② 나 ③ 다 ④ 라 ⑤ 마\n2. 둘째 문항입니다.\n① 가 ② 나 ③ 다 ④ 라 ⑤ 마\n"


def check(condition: bool, message: str) -> None:
    if not condition:
        failures.append(message)


def rows(result: dict) -> list[dict]:
    return [*result.get("created", []), *result.get("existing", [])]


def expect_two_questions(label: str, result: dict, answers: tuple[str, str] = ("③", "⑤")) -> None:
    items = rows(result)
    numbers = [item["number"] for item in items]
    check(numbers == ["1", "2"], f"{label}: answer key became questions, numbers={numbers}")
    check([item["answer"] for item in items] == list(answers),
          f"{label}: answers {[item['answer'] for item in items]}, expected {list(answers)}")
    check(all("정답" not in item["stem"] for item in items), f"{label}: answer-key header leaked into a stem")
    check(any("정답·해설 부분은 문항으로 만들지 않았습니다" in notice for notice in result.get("notices", [])),
          f"{label}: missing answer-key notice: {result.get('notices')}")


def text_and_docx_checks() -> None:
    text = QUESTIONS + "빠른 정답\n1. ③ 2. ⑤\n"
    expect_two_questions("TXT", importers.import_text("key.txt", text.encode("utf-8"), {}))

    document = Document()
    for line in text.strip().splitlines():
        document.add_paragraph(line)
    buffer = io.BytesIO()
    document.save(buffer)
    expect_two_questions("DOCX", importers.import_docx("key.docx", buffer.getvalue(), {}))

    grid = Document()
    for line in QUESTIONS.strip().splitlines():
        grid.add_paragraph(line)
    grid.add_paragraph("정답표")
    table = grid.add_table(rows=2, cols=2)
    for column, (number, answer) in enumerate((("1", "②"), ("2", "④"))):
        table.cell(0, column).text = number
        table.cell(1, column).text = answer
    buffer = io.BytesIO()
    grid.save(buffer)
    expect_two_questions("DOCX grid", importers.import_docx("grid.docx", buffer.getvalue(), {}), ("②", "④"))

    # Per-question answer lines are not an answer section; nothing may be cut.
    inline = "1. 첫째 문항입니다.\n[정답] ③\n2. 둘째 문항입니다.\n[정답] ⑤\n3. 셋째 문항입니다.\n"
    inline_rows = rows(importers.import_text("inline.txt", inline.encode("utf-8"), {}))
    check([item["number"] for item in inline_rows] == ["1", "2", "3"], f"inline [정답] lines cut questions: {inline_rows}")
    check(all("③" in (item["stem"] + item["answer"]) for item in inline_rows[:1]), f"inline answer line lost: {inline_rows}")

    # An explanation booklet (header before any question) keeps its old behaviour.
    booklet = "정답 및 해설\n1. ③ 첫 문항 해설\n2. ⑤ 둘째 문항 해설\n"
    booklet_rows = rows(importers.import_text("booklet.txt", booklet.encode("utf-8"), {}))
    check(len(booklet_rows) == 2, f"explanation booklet lost its items: {booklet_rows}")

    # Real answer pages put the exam title before the header ('… 화학I 정답 및 해설', '[정답표]').
    titled = QUESTIONS + "2025학년도 3월 모의고사 정답 및 해설\n[정답표] 1. ③ 2. ⑤\n1. [해설] 정답 ③ 해설입니다.\n"
    expect_two_questions("TXT titled header", importers.import_text("titled.txt", titled.encode("utf-8"), {}))
    # A question stem that merely ends with '정답과 해설' (no answer pairs after it) is not cut.
    stem_like = QUESTIONS + "3. 다음은 학생이 정리한 정답과 해설\n① 가 ② 나 ③ 다 ④ 라 ⑤ 마\n"
    stem_rows = rows(importers.import_text("stem_like.txt", stem_like.encode("utf-8"), {}))
    check([item["number"] for item in stem_rows] == ["1", "2", "3"], f"stem ending in 정답과 해설 was cut: {stem_rows}")

    # Unit-by-unit answer keys: questions that continue with the next number after a key
    # are real questions and must survive (the old cut dropped everything after the key).
    units = (QUESTIONS + "빠른 정답\n1. ③ 2. ⑤\n"
             "3. 셋째 문항입니다.\n① 가 ② 나 ③ 다 ④ 라 ⑤ 마\n빠른 정답\n3. ①\n")
    unit_rows = rows(importers.import_text("units.txt", units.encode("utf-8"), {}))
    check([item["number"] for item in unit_rows] == ["1", "2", "3"], f"questions after a unit answer key were lost: {unit_rows}")
    check([item["answer"] for item in unit_rows] == ["③", "⑤", "①"], f"unit answer keys not mapped: {unit_rows}")

    # Without a matching question number the answer is reported, not invented.
    stray = QUESTIONS + "빠른 정답\n1. ③ 7. ①\n"
    stray_result = importers.import_text("stray.txt", stray.encode("utf-8"), {})
    check(any("넣지 못했습니다" in notice for notice in stray_result.get("notices", [])),
          f"unmatched answer not reported: {stray_result.get('notices')}")


def hwpx_endnote_fallback_check() -> None:
    hp = "http://www.hancom.co.kr/hwpml/2011/paragraph"

    def para(text: str, note: str = "", number: str = "") -> str:
        endnote = (
            f'<hp:ctrl><hp:endNote number="{number}"><hp:subList><hp:p><hp:run><hp:t>{note}</hp:t>'
            f"</hp:run></hp:p></hp:subList></hp:endNote></hp:ctrl>"
        ) if note else ""
        return f"<hp:p><hp:run><hp:t>{text}</hp:t>{endnote}</hp:run></hp:p>"

    # Exam starts at 21 and the notes carry no '[답]' lines: endnote boundaries do not apply.
    section = (
        f'<hs:sec xmlns:hs="http://www.hancom.co.kr/hwpml/2011/section" xmlns:hp="{hp}">'
        + para("21. 첫 문항입니다.", "첫 문항 해설입니다.", "1")
        + para("① 가 ② 나 ③ 다 ④ 라 ⑤ 마")
        + para("22. 둘째 문항입니다.", "둘째 문항 해설입니다.", "2")
        + para("① 가 ② 나 ③ 다 ④ 라 ⑤ 마")
        + "</hs:sec>"
    )
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("mimetype", "application/hwp+zip")
        archive.writestr("Contents/section0.xml", section)
    result = importers.import_hwpx("endnote.hwpx", buffer.getvalue(), {})
    items = rows(result)
    check([item["number"] for item in items] == ["21", "22"], f"HWPX endnote fallback numbers: {items}")
    check([item["explanation"] for item in items] == ["첫 문항 해설입니다.", "둘째 문항 해설입니다."],
          f"HWPX endnote explanations were dropped: {[item['explanation'] for item in items]}")
    check(any("미주에 있던 정답·해설 2개 중 2개" in notice for notice in result.get("notices", [])),
          f"HWPX endnote fallback notice missing: {result.get('notices')}")


def pdf_answer_page_check(title: str = "") -> None:
    document = fitz.open()
    page = document.new_page()
    font = {"fontname": "korea", "fontsize": 11}
    y = 80
    for number in (1, 2, 3):
        page.insert_text((60, y), f"{number}. 다음 중 옳은 것은? (합성 문항 {number})", **font)
        page.insert_text((60, y + 22), "① 가  ② 나  ③ 다  ④ 라  ⑤ 마", **font)
        y += 90
    key = document.new_page()
    # title: real answer pages start with the exam title, e.g. '2025학년도 … 화학I 정답 및 해설'.
    key.insert_text((60, 80), f"{title}정답 및 해설", **font)
    for index, answer in enumerate(("②", "④", "①"), start=1):
        key.insert_text((60, 80 + 24 * index), f"{index}번 정답 {answer}  해설: {index}번은 선택지를 비교한다.", **font)
    label = f"PDF{' titled' if title else ''}"
    result = importers.import_pdf(f"key{len(title)}.pdf", document.tobytes(), {"math_ai_recognition": False})
    items = rows(result)
    numbers = [item["number"] for item in items]
    check(numbers == ["1", "2", "3"], f"{label} answer page became questions: numbers={numbers}")
    check([item["answer"] for item in items] == ["②", "④", "①"], f"{label} answers not mapped: {[item['answer'] for item in items]}")
    check(all(int(item.get("source_page") or 0) == 1 for item in items), f"{label} kept an item from the answer page")


def pdf_unit_answer_page_check() -> None:
    """Questions 3-4 on a page after the answer page for 1-2 are real questions."""
    document = fitz.open()
    font = {"fontname": "korea", "fontsize": 11}
    for numbers, answers in (((1, 2), ("②", "④")), ((3, 4), ("①", "③"))):
        page = document.new_page()
        y = 80
        for number in numbers:
            page.insert_text((60, y), f"{number}. 다음 중 옳은 것은? (합성 문항 {number})", **font)
            page.insert_text((60, y + 22), "① 가  ② 나  ③ 다  ④ 라  ⑤ 마", **font)
            y += 90
        key = document.new_page()
        key.insert_text((60, 80), "정답 및 해설", **font)
        for index, (number, answer) in enumerate(zip(numbers, answers), start=1):
            key.insert_text((60, 80 + 24 * index), f"{number}번 정답 {answer}  해설: 선택지를 비교한다.", **font)
    result = importers.import_pdf("units.pdf", document.tobytes(), {"math_ai_recognition": False})
    items = rows(result)
    numbers = [item["number"] for item in items]
    check(numbers == ["1", "2", "3", "4"], f"PDF unit answer pages: numbers={numbers}")
    check([item["answer"] for item in items][:2] == ["②", "④"], f"PDF unit answers: {[item['answer'] for item in items]}")


def inline_answer_regression_checks() -> None:
    """Questions that carry their own '[정답] ③' lines are never an answer section (repair R1).

    Before the fix: F1 PDF page 2 was taken as an answer page (made-up answer ① for 4-6 and a
    false notice), F2 lost questions 3 and 6, F3/F3b read the numeric choices '① 1 ② 2 …' as
    answer pairs (choices lost, made-up answers ②③④, false notice).
    """
    def no_key_notice(label: str, result: dict) -> None:
        check(not any("정답·해설 부분" in notice for notice in result.get("notices", [])),
              f"{label}: false answer-key notice {result.get('notices')}")

    # F1: 2-page PDF, 3 questions per page, each followed by an inline '[정답] ③' line.
    document = fitz.open()
    font = {"fontname": "korea", "fontsize": 11}
    number = 1
    for _page in range(2):
        page = document.new_page()
        y = 80
        for _question in range(3):
            page.insert_text((60, y), f"{number}. 다음 중 옳은 것은? (합성 문항 {number})", **font)
            page.insert_text((60, y + 22), "① 가  ② 나  ③ 다  ④ 라  ⑤ 마", **font)
            page.insert_text((60, y + 44), "[정답] ③", **font)
            y += 110
            number += 1
    result = importers.import_pdf("inline_key.pdf", document.tobytes(), {"math_ai_recognition": False})
    items = rows(result)
    check([item["number"] for item in items] == ["1", "2", "3", "4", "5", "6"],
          f"F1 PDF inline [정답]: numbers={[item['number'] for item in items]}")
    check(all(item["answer"] in ("", "③") for item in items),
          f"F1 PDF inline [정답]: made-up answers {[item['answer'] for item in items]}")
    no_key_notice("F1 PDF inline [정답]", result)

    # F2: one DOCX paragraph holds questions 3-5 (soft breaks) with inline answers.
    doc = Document()
    for line in ("1. 첫째 문항입니다.", "① 가 ② 나 ③ 다 ④ 라 ⑤ 마", "2. 둘째 문항입니다.", "① 가 ② 나 ③ 다 ④ 라 ⑤ 마"):
        doc.add_paragraph(line)
    paragraph = doc.add_paragraph()
    pieces = ["3. 다음 중 옳은 것은?", "① 가 ② 나 ③ 다 ④ 라 ⑤ 마", "[정답] ③", "4. 넷째 문항은?", "[정답] ①",
              "5. 다섯째 문항은?", "[정답] ②"]
    for index, piece in enumerate(pieces):
        run = paragraph.add_run(piece)
        if index < len(pieces) - 1:
            run.add_break()
    doc.add_paragraph("6. 여섯째 문항입니다.")
    doc.add_paragraph("① 가 ② 나 ③ 다 ④ 라 ⑤ 마")
    buffer = io.BytesIO()
    doc.save(buffer)
    result = importers.import_docx("soft_inline.docx", buffer.getvalue(), {})
    numbers = [item["number"] for item in rows(result)]
    check("3" in numbers and "6" in numbers, f"F2 DOCX soft-broken inline [정답]: questions lost, numbers={numbers}")
    no_key_notice("F2 DOCX soft-broken inline [정답]", result)
    check(not importers._answer_key_start("3. 다음 중 옳은 것은?\n[정답] ③\n4. 넷째\n[정답] ①\n5. 다섯째\n[정답] ②"),
          "a block that starts with a question number is not an answer-section start")

    # F3/F3b: a stem line ending in '정답표'/'빠른 정답' followed by numeric choices.
    check(importers._parse_answer_key(["① 1 ② 2 ③ 3 ④ 4 ⑤ 5"]) == {},
          "numeric choice list is not read as answer pairs")
    check(importers._parse_answer_key(["[정답표] 1. ③  2. ①  3. ⑤"]) == {"1": "③", "2": "①", "3": "⑤"},
          "a real answer table after circled answers is still read")
    for label, header in (("F3", "철수가 만든 정답표"), ("F3b", "영희의 빠른 정답")):
        text = (f"1. 채점 결과 틀린 문항 수는?\n{header}\n① 1 ② 2 ③ 3 ④ 4 ⑤ 5\n"
                "2. 둘째 문항은?\n① 가 ② 나 ③ 다 ④ 라 ⑤ 마\n3. 셋째 문항은?\n① 가 ② 나 ③ 다 ④ 라 ⑤ 마\n")
        result = importers.import_text(f"{label}.txt", text.encode("utf-8"), {})
        items = rows(result)
        check([item["number"] for item in items] == ["1", "2", "3"], f"{label}: numbers={[item['number'] for item in items]}")
        check(all(not item["answer"] for item in items), f"{label}: made-up answers {[item['answer'] for item in items]}")
        check(bool(items) and len(items[0].get("choices") or []) == 5 and header in items[0]["stem"],
              f"{label}: question 1 lost its line or choices: {items[:1]}")
        no_key_notice(label, result)


def quick_answer_and_official_key_checks() -> None:
    """Space/line separated quick answers and the official CSAT key layout (repair R1, verifier F4/2e).

    Before the fix the numeric-choice guard kept only the first pair of '1 ③ 2 ① 3 ④',
    '1 ③ / 2 ①' (line separated) and '01 ③ 02 ①', and a title header ('수학 영역 정답표') checked
    its lines one by one, so the official key (number and circled answer on separate lines)
    was never an answer page (36/36 -> 0/36).
    """
    nl = "\n"
    expected = {"1": "③", "2": "①", "3": "④"}
    for sample in ("1 ③ 2 ① 3 ④", nl.join(["1 ③", "2 ①", "3 ④"]), "01 ③ 02 ① 03 ④", "1. ③ 2. ① 3. ④",
                   "1번 ③ 2번 ① 3번 ④"):
        got = importers._parse_answer_key([sample])
        check(got == expected, f"quick answers {sample!r}: {got}")
    # answers that happen to follow the choice order are still answers, not a choice list
    check(importers._parse_answer_key(["1 ② 2 ③ 3 ④"]) == {"1": "②", "2": "③", "3": "④"},
          "quick answers '1 ② 2 ③ 3 ④' must keep all pairs")
    check(importers._parse_answer_key(["① 1 ② 2 ③ 3 ④ 4 ⑤ 5"]) == {}, "numeric choice list read as pairs")
    check(importers._parse_answer_key(["[정답표] 1. ③  2. ①"]) == {"1": "③", "2": "①"}, "[정답표] pairs lost")
    official = nl.join(["2026학년도 대학수학능력시험", "한국사 영역 정답표", "( 홀수 ) 형", "문항", "번호", "정 답배 점",
                        "1", "③", "2", "6", "⑤", "2", "11", "③", "3", "2", "③", "2", "7", "④", "3"])
    check(importers._pdf_page_answer_key_start(official),
          "official CSAT key layout (number/answer on separate lines) not an answer page")
    check(not importers._pdf_page_answer_key_start(nl.join(["1. 철수가 채점한 결과는?", "철수가 만든 정답표",
                                                            "① 1 ② 2 ③ 3 ④ 4 ⑤ 5"])),
          "F3 stem with numeric choices became an answer page")

    text = nl.join(["1. 첫째 문항입니다.", "① 가 ② 나 ③ 다 ④ 라 ⑤ 마", "2. 둘째 문항입니다.", "① 가 ② 나 ③ 다 ④ 라 ⑤ 마",
                    "3. 셋째 문항입니다.", "① 가 ② 나 ③ 다 ④ 라 ⑤ 마", "빠른 정답", "1 ③ 2 ① 3 ④", ""])
    result = importers.import_text("quick_space.txt", text.encode("utf-8"), {})
    items = sorted(rows(result), key=lambda item: int(item["number"] or 0))
    check([item["number"] for item in items] == ["1", "2", "3"], f"TXT quick space answers: numbers={[item['number'] for item in items]}")
    check([item["answer"] for item in items] == ["③", "①", "④"], f"TXT quick space answers: {[item['answer'] for item in items]}")

    # Real official key (data/ original). Only this part SKIPs when the file is absent.
    csat = ROOT / "data" / "external_exam_qa" / "2026_csat"
    keys = sorted((csat / "정답표").glob("**/*.pdf"))
    history_key = csat / "정답표" / "한국사" / "한국사.pdf"
    if not history_key.exists():
        print("  SKIP: real CSAT answer-key PDFs not found (official key checks only)")
        return
    starts = sum(importers._pdf_page_answer_key_start(fitz.open(path)[0].get_text()) for path in keys)
    check(starts == len(keys), f"real CSAT answer keys detected as answer pages: {starts}/{len(keys)}")
    history_exam = sorted((csat / "문제지" / "한국사").glob("*.pdf"))
    if history_exam:
        combined = fitz.open(history_exam[0])
        combined.insert_pdf(fitz.open(history_key))
        result = importers.import_pdf("history_with_key.pdf", combined.tobytes(), {"math_ai_recognition": False})
        items = rows(result)
        check(all(item["number"] for item in items), f"history+key PDF: unnumbered fake questions {[item['number'] for item in items]}")
        filled = sum(1 for item in items if item["answer"])
        check(filled >= 20, f"history+key PDF: only {filled} answers filled")


def collect_check(client: TestClient) -> None:
    site = Path(_TMP.name) / "site"
    site.mkdir()
    (site / "exam.html").write_text(
        "<!doctype html><html><head><meta charset='utf-8'><title>수집 점검</title></head><body><article>"
        "<p>1. 다음 중 가장 큰 수는?</p><p>① 0.5 ② 1.5 ③ 2 ④ 3.5 ⑤ 12.5</p>"
        "<p>2. 다음 중 단위가 맞는 것은?</p><p>① 2 cm ② 3 kg ③ 4 L ④ 5 m ⑤ 6 g</p>"
        "<p>빠른 정답</p><p>1. ⑤ 2. ①</p></article></body></html>",
        encoding="utf-8",
    )

    class QuietHandler(http.server.SimpleHTTPRequestHandler):
        def log_message(self, *_args: object) -> None:
            pass

    server = http.server.ThreadingHTTPServer(("127.0.0.1", 0), functools.partial(QuietHandler, directory=str(site)))
    threading.Thread(target=server.serve_forever, daemon=True).start()
    original = collector._validate_public_host
    try:
        # Loopback fetch for this in-process test only; the SSRF guard itself is pinned in verify_api_hardening.
        collector._validate_public_host = lambda hostname, port=None: None
        response = client.post("/api/collect", json={"url": f"http://127.0.0.1:{server.server_address[1]}/exam.html"})
    finally:
        collector._validate_public_host = original
        server.shutdown()
        server.server_close()
    check(response.status_code == 200, f"/api/collect returned {response.status_code}: {response.text[:200]}")
    if response.status_code == 200:
        items = rows(response.json())
        check([item["number"] for item in items] == ["1", "2"], f"collect numbers: {[item['number'] for item in items]}")
        check([item["answer"] for item in items] == ["⑤", "①"], f"collect answers: {[item['answer'] for item in items]}")


def main_check() -> int:
    storage.init_db()
    text_and_docx_checks()
    hwpx_endnote_fallback_check()
    pdf_answer_page_check()
    pdf_answer_page_check("2025학년도 대학수학능력시험 과학탐구영역 화학I ")
    pdf_unit_answer_page_check()
    inline_answer_regression_checks()
    quick_answer_and_official_key_checks()
    with TestClient(main.app) as client:
        collect_check(client)
    if failures:
        print(f"Answer key import verification FAILED ({len(failures)} issues)")
        for failure in failures:
            print(f"  - {failure}")
        return 1
    print("Answer key import verification PASS")
    print("  TXT/DOCX/DOCX grid/web/PDF trailing answer key -> answers, no fake questions: PASS")
    print("  answer-key header after an exam title ('… 정답 및 해설', '[정답표]') in TXT/PDF: PASS")
    print("  per-question [정답] lines and explanation booklets unchanged: PASS")
    print("  questions after a unit-by-unit answer key (TXT/PDF) are kept: PASS")
    print("  inline [정답] pages/paragraphs and '정답표' stems with numeric choices are not cut: PASS")
    print("  quick answers '1 ③ 2 ①' / line-separated and the official CSAT key layout: PASS")
    print("  HWPX endnotes outside the boundary rule attach instead of vanishing: PASS")
    print("  /api/collect on a normal page: PASS")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main_check())
    finally:
        try:
            with storage.connect() as connection:
                connection.execute("PRAGMA wal_checkpoint(TRUNCATE)")
                connection.execute("PRAGMA journal_mode = DELETE")
        except Exception:
            pass
        _TMP.cleanup()
