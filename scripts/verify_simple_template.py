"""회귀 테스트: 간단 변환 전용 'simple' 내보내기 양식(2026-10-03 4단계).

간단 변환·Windows 앱·웹 체험판의 기본 변환은 basic(문항마다 '정답:/해설:', 'N. 파일명 #N'
제목 줄, '1)' 선지) 대신 simple 양식을 쓴다. 합성 미주형 문항 세트(정답·해설 포함)로
HWPX(v2·구 v1)·DOCX writer 의 불변식을 고정한다.

- 문항 아래 '정답:'/'해설:' 줄이 없고, 정답은 문서 끝 새 쪽 '정답' 표에만 있다.
- '#N' 제목 줄이 없고, 원번호가 문항마다 정확히 한 번(원본과 같은 순서) 나온다.
- 선지는 원문자(①~⑤)로 표기한다.
- 정답이 없는 입력엔 정답표를 붙이지 않고, 옵션으로 끌 수 있고, 전체 정답지와 겹치지 않는다.
- 지문 묶음 항목(번호 '[1~3]')엔 번호를 덧붙이지 않는다('[1~3]. [1~3] …' 중복 금지).
- 발문 뒤 '1) …'·'2) …' 조건 줄은 번호 줄로 바꾸지 않는다(번호는 발문 첫 줄에).
- basic 양식 출력(정답·해설·'#N'·'1)')은 그대로다.
- 목록에는 포함되지만 스튜디오 기본 양식(목록 첫 항목)은 basic 이다.
- 간단 변환·데스크톱·웹 체험판이 simple 을 지정한다.

개인 샘플 무의존. 종료코드 0=통과, 1=실패. scripts/run_all_verify.py 자동 편입.
"""
from __future__ import annotations

import gc
import os
import re
import sys
import tempfile
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

_TMP = tempfile.TemporaryDirectory(prefix="hwp_make_simple_template_")
os.environ["HWP_MAKE_DATA_DIR"] = _TMP.name

from docx import Document  # noqa: E402
from lxml import etree  # noqa: E402

from app import docx_writer, exam_templates, hwpx_writer, hwpx_writer_v2  # noqa: E402

HP = "{http://www.hancom.co.kr/hwpml/2011/paragraph}"
CIRCLED = "①②③④⑤"
TITLE = "2025 미주형 모의고사"

# 미주형 HWPX 가져오기 결과와 같은 모양: 번호는 number 에만 있거나 본문 첫 줄에도 있고,
# 묶음 지문은 첫 문항 본문 앞에 붙어 있으며, 정답·해설이 문항마다 있다.
PROBLEMS = [
    {"number": "1", "title": "모의 #1",
     "stem": "[1~2] 다음 글을 읽고 물음에 답하시오.\n지문 첫 문단이다.\n지문 둘째 문단이다.\n윗글의 내용과 일치하지 않는 것은?",
     "choices": ["가", "나", "다", "라", "마"], "answer": "③", "explanation": "지문 해설"},
    {"number": "2", "title": "모의 #2", "stem": "2. 윗글을 바탕으로 이해한 것으로 적절한 것은? [3점]",
     "choices": ["① 하나", "② 둘", "③ 셋", "④ 넷", "⑤ 다섯"], "answer": "5", "explanation": "둘째 해설"},
    {"number": "3", "title": "모의 #3", "stem": "합성 시험지 머리 줄\n3. 3.5보다 큰 수는?",
     "choices": ["1) 1", "2) 2", "3) 3", "4) 4", "5) 5"], "answer": "4", "explanation": ""},
    {"number": "4", "title": "모의 #4", "stem": "4) 단답형: 값을 구하시오.", "choices": [],
     "answer": "12", "explanation": "단답 해설"},
]
NUMBERS = ["1", "2", "3", "4"]
# 국어 HWP 가져오기의 지문 묶음 안내 줄 모양(범위 양옆 U+2007 공백).
PASSAGE_HEAD = "[1\u2007~\u20072] 다음 글을 읽고 물음에 답하시오."

failures: list[str] = []


def check(condition: bool, message: str) -> None:
    print(f"  [{'PASS' if condition else 'FAIL'}] {message}")
    if not condition:
        failures.append(message)


def hwpx_paragraphs(path: Path) -> tuple[list[str], list[list[str]], list[bool]]:
    """(본문 최상위 문단, 표 행 텍스트 목록, 문단별 쪽 나눔) — 표 안 문단은 따로 모은다."""
    with zipfile.ZipFile(path) as archive:
        root = etree.fromstring(archive.read("Contents/section0.xml"))
    top = root.findall(HP + "p")

    def own_text(paragraph) -> str:
        parts = []
        for run in paragraph.findall(HP + "run"):
            parts.extend("".join(t.itertext()) for t in run.findall(HP + "t"))
        return "".join(parts)

    rows = [["".join("".join(t.itertext()) for t in cell.iter(HP + "t")) for cell in tr.findall(HP + "tc")]
            for tr in root.iter(HP + "tr")]
    return [own_text(p) for p in top], rows, [p.get("pageBreak") == "1" for p in top]


def docx_paragraphs(path: Path) -> tuple[list[str], list[list[str]]]:
    document = Document(str(path))
    rows = [[cell.text for cell in row.cells] for table in document.tables for row in table.rows]
    return [p.text for p in document.paragraphs], rows


def heading_numbers(lines: list[str]) -> list[str]:
    return [m.group(1) for m in (re.match(r"^\s*(\d{1,3})\s*[.)]", line) for line in lines) if m]


def simple_invariants(kind: str, lines: list[str], rows: list[list[str]]) -> None:
    check(not any(line.startswith(("정답:", "해설:")) for line in lines), f"{kind}: 문항 아래 정답·해설 줄 없음")
    check(not any("해설" in line for line in lines), f"{kind}: 해설 본문은 담지 않음")
    check(not any(re.search(r"#\d+", line) for line in lines), f"{kind}: '#N' 제목 줄 없음")
    check(heading_numbers(lines) == NUMBERS, f"{kind}: 원번호가 한 번씩 순서대로 {heading_numbers(lines)}")
    choice_lines = [line for line in lines if line and line[0] in CIRCLED]
    check(len(choice_lines) == 15 and not any(re.match(r"^[1-5]\)\s", line) for line in lines),
          f"{kind}: 선지는 원문자 표기({len(choice_lines)}줄)")
    check("2. 윗글을 바탕으로 이해한 것으로 적절한 것은? [3점]" in lines, f"{kind}: 본문 번호와 겹치지 않음")
    check("3. 3.5보다 큰 수는?" in lines, f"{kind}: 번호 줄 앞의 머리 줄은 본문으로 두고 번호는 한 번")
    check("1. 윗글의 내용과 일치하지 않는 것은?" in lines
          and lines.index("[1~2] 다음 글을 읽고 물음에 답하시오.") < lines.index("1. 윗글의 내용과 일치하지 않는 것은?"),
          f"{kind}: 묶음 지문 뒤 발문 줄에 번호")
    check("정답" in lines, f"{kind}: 정답 제목 존재")
    check(lines.index("정답") > lines.index("4. 단답형: 값을 구하시오."), f"{kind}: 정답표는 마지막 문항 뒤")
    check(rows[-2:] == [["번호", *NUMBERS], ["정답", "③", "⑤", "④", "12"]], f"{kind}: 정답표 {rows[-2:]}")


def main() -> int:
    template = exam_templates.TEMPLATE_MAP.get("simple")
    check(template is not None and template.label == "기본 변환", "simple 양식 등록·라벨 '기본 변환'")
    if template is None:
        print("SIMPLE_TEMPLATE_FAIL")
        return 1
    check(exam_templates.TEMPLATES[0].key == "basic", "스튜디오 기본(목록 첫 항목)은 basic 유지")
    check((template.include_answers, template.include_explanations, template.circled_choices,
           template.native_math_default, template.answer_key_appendix, template.source_number_line)
          == (False, False, True, True, True, True), "simple 필드 값")
    check(exam_templates.get_template(None).key == "basic", "양식 미지정 기본값은 basic")
    check(exam_templates.resolve_export_title(TITLE, template) == TITLE, "문서 제목은 넘긴 파일명 그대로")

    # 번호 줄은 앞 세 줄 안에서만 찾는다(본문 깊은 곳의 'N.' 줄을 번호 줄로 오인하지 않게).
    far = exam_templates.numbered_stem_paragraphs(["머리 줄 가", "머리 줄 나", "머리 줄 다", "7. 넷째 줄"], "7")
    check(far[0] == ("7. 머리 줄 가", "heading"), f"넷째 줄 이후는 번호 줄로 보지 않음 {far[0]}")

    with tempfile.TemporaryDirectory(prefix="hwp_make_simple_out_") as tmp:
        out = Path(tmp)
        hwpx_writer_v2.write_hwpx(out / "simple.hwpx", TITLE, PROBLEMS, "simple", native_math=True)
        lines, rows, breaks = hwpx_paragraphs(out / "simple.hwpx")
        check(TITLE in lines[:3], "HWPX: 제목은 파일명 일반 텍스트")
        simple_invariants("HWPX", lines, rows)
        check(breaks[lines.index("정답")], "HWPX: 정답표는 새 쪽에서 시작")

        # 구 writer(v1)도 같은 의미여야 coverage_hwpx_v2 의 v1/v2 본문 패리티가 유지된다.
        hwpx_writer.write_hwpx(out / "simple_v1.hwpx", TITLE, PROBLEMS, "simple")
        v1_lines, v1_rows, v1_breaks = hwpx_paragraphs(out / "simple_v1.hwpx")
        check(TITLE in v1_lines[:3], "HWPX v1: 제목은 파일명 일반 텍스트")
        simple_invariants("HWPX v1", v1_lines, v1_rows)
        check(v1_breaks[v1_lines.index("정답")], "HWPX v1: 정답표는 새 쪽에서 시작")

        docx_writer.write_docx(out / "simple.docx", TITLE, PROBLEMS, "simple")
        dlines, drows = docx_paragraphs(out / "simple.docx")
        check(dlines[0] == TITLE, "DOCX: 제목은 파일명 일반 텍스트")
        simple_invariants("DOCX", dlines, drows)

        no_answers = [{**problem, "answer": "", "explanation": ""} for problem in PROBLEMS]
        hwpx_writer_v2.write_hwpx(out / "no_answers.hwpx", TITLE, no_answers, "simple")
        lines, rows, _ = hwpx_paragraphs(out / "no_answers.hwpx")
        check("정답" not in lines and not rows, "정답이 없는 입력엔 정답표 없음")

        # 지문 묶음 항목(번호 '[1~3]', 국어 HWP 가져오기 모양)은 번호를 덧붙이지 않는다.
        # 수정 전 DOCX·v1 은 '[1~3]. [1~3] 다음 글을…' 처럼 범위를 두 번 썼다.
        passage_set = [
            {"number": "[1~2]", "title": "지문[1~2]",
             "stem": PASSAGE_HEAD + "\n지문 본문 줄이다.", "choices": []},
            {**PROBLEMS[1], "stem": "1. 윗글의 내용으로 적절한 것은?", "number": "1"},
        ]
        for kind, writer, suffix in (("HWPX", hwpx_writer_v2.write_hwpx, "hwpx"),
                                     ("HWPX v1", hwpx_writer.write_hwpx, "hwpx"),
                                     ("DOCX", docx_writer.write_docx, "docx")):
            target = out / f"passage_{kind.replace(' ', '_')}.{suffix}"
            writer(target, TITLE, passage_set, "simple")
            plines = (hwpx_paragraphs(target)[0] if suffix == "hwpx" else docx_paragraphs(target)[0])
            check(PASSAGE_HEAD in plines and not any(line.startswith("[1~2].") for line in plines),
                  f"{kind}: 지문 묶음 안내 줄에 번호를 덧붙이지 않음")
            check("지문 본문 줄이다." in plines and "1. 윗글의 내용으로 적절한 것은?" in plines,
                  f"{kind}: 지문 본문과 뒤 문항 번호 유지")
        # 둘째 줄부터의 '1)'·'2)' 조건 줄은 번호 줄이 아니다(수정 전: '2) f(1)=2' 가 '2. f(1)=2' 번호 줄이 됨).
        direct = exam_templates.numbered_stem_paragraphs(
            ["함수 f(x)가 다음 조건을 만족시킨다.", "1) f(0)=1", "2) f(1)=2"], "2")
        check(direct == [("2. 함수 f(x)가 다음 조건을 만족시킨다.", "heading"),
                         ("1) f(0)=1", "body"), ("2) f(1)=2", "body")],
              f"조건 줄 '1)'·'2)' 는 번호 줄로 쓰지 않음 {direct}")
        # 수식 없는 조건 줄: DOCX 는 f(x) 같은 식을 수식 개체로 내보내 문단 텍스트에서 빠진다.
        condition_set = [{"number": "2", "title": "조건",
                          "stem": "다음 조건을 모두 만족시키는 수는?\n1) 짝수이다.\n2) 한 자리 수이다.",
                          "choices": ["1", "2", "3", "4", "5"], "answer": "3"}]
        for kind, writer, suffix in (("HWPX", hwpx_writer_v2.write_hwpx, "hwpx"),
                                     ("HWPX v1", hwpx_writer.write_hwpx, "hwpx"),
                                     ("DOCX", docx_writer.write_docx, "docx")):
            target = out / f"condition_{kind.replace(' ', '_')}.{suffix}"
            writer(target, TITLE, condition_set, "simple")
            clines = (hwpx_paragraphs(target)[0] if suffix == "hwpx" else docx_paragraphs(target)[0])
            check("2. 다음 조건을 모두 만족시키는 수는?" in clines
                  and "1) 짝수이다." in clines and "2) 한 자리 수이다." in clines
                  and not any(line.startswith("2. 한 자리") for line in clines),
                  f"{kind}: 번호는 발문 첫 줄에, 조건 줄은 원문 그대로")
        hwpx_writer_v2.write_hwpx(out / "off.hwpx", TITLE, PROBLEMS, "simple", answer_key_appendix=False)
        lines, rows, _ = hwpx_paragraphs(out / "off.hwpx")
        check("정답" not in lines and not rows, "정답표 옵션 끄기")
        hwpx_writer_v2.write_hwpx(out / "sheet.hwpx", TITLE, PROBLEMS, "simple", include_answer_sheet=True)
        lines, rows, _ = hwpx_paragraphs(out / "sheet.hwpx")
        check(exam_templates.ANSWER_SHEET_TITLE in lines and "정답" not in lines and not rows,
              "전체 정답지를 붙이면 정답표는 중복하지 않음")

        # basic 은 기존 동작 그대로여야 한다(정답·해설 줄, '#N' 제목, '1)' 선지, 끝 정답표 없음).
        hwpx_writer_v2.write_hwpx(out / "basic.hwpx", TITLE, PROBLEMS, "basic")
        lines, rows, _ = hwpx_paragraphs(out / "basic.hwpx")
        check(sum(line.startswith("정답:") for line in lines) == 4
              and sum(line.startswith("해설:") for line in lines) == 3, "basic: 문항별 정답·해설 유지")
        check(sum(bool(re.search(r"#\d+$", line)) for line in lines) == 4, "basic: 'N. 제목 #N' 줄 유지")
        check(sum(bool(re.match(r"^[1-5]\) ", line)) for line in lines) == 15 and "정답" not in lines and not rows,
              "basic: '1)' 선지 유지, 끝 정답표 없음")
        docx_writer.write_docx(out / "basic.docx", TITLE, PROBLEMS, "basic")
        dlines, drows = docx_paragraphs(out / "basic.docx")
        check(sum(line.startswith("정답:") for line in dlines) == 4 and not drows, "basic DOCX: 정답 줄 유지, 정답표 없음")

    from app import main as app_main  # noqa: E402 - 격리 데이터 폴더 지정 뒤에 불러온다.

    keys = [item["key"] for item in app_main.export_templates()["items"]]
    check("simple" in keys and keys[0] == "basic", f"/api/export-templates 에 simple 포함, 첫 항목 basic ({keys[:3]})")
    check(app_main.ExportPayload(ids=[1]).answer_key_appendix is None, "API 정답표 옵션 기본은 양식 기본값")

    wiring = {
        "app/desktop_convert.py": 'template_key="simple"',
        "app/web_convert.py": 'template_key="simple"',
        "static/app.js": 'const SIMPLE_EXPORT_TEMPLATE = "simple";',
    }
    for relative, needle in wiring.items():
        source = (ROOT / relative).read_text(encoding="utf-8")
        check(needle in source, f"{relative}: 기본 변환이 simple 양식 지정")

    if failures:
        print(f"SIMPLE_TEMPLATE_FAIL ({len(failures)}건)")
        return 1
    print("SIMPLE_TEMPLATE_OK")
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    finally:
        gc.collect()  # storage 의 sqlite 연결을 닫아야 Windows 에서 임시 폴더가 지워진다.
        _TMP.cleanup()
